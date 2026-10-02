"""
Gera o modelo Word (docxtpl) do Relatório Anual a partir do relatório de 2025.

Ferramenta de desenvolvimento — não é usada em produção. O resultado
(modelo_relatorio_anual.docx) é versionado no repositório.

    pip install docxtpl
    python portal/catequista/relatorio_anual/build_modelo.py "docs/Relatório 2025.docx"

O modelo mantém o cabeçalho (logótipos), estilos, numeração e tabela do
original; o corpo é substituído por marcadores Jinja que o docxtpl preenche.
O coordenador pode depois editar o .docx no Word (mover secções, mudar
fontes) desde que não parta as etiquetas {{ ... }} / {%p ... %}.
"""

import copy
import os
import sys

from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "modelo_relatorio_anual.docx")

# Índices dos parágrafos exemplares no Relatório 2025.docx
EX_TITULO = 0       # centrado, negrito, 12pt
EX_LEMA = 5         # centrado
EX_CORPO = 6        # justificado
EX_MES = 8          # negrito, justificado
EX_MARCADOR = 9     # List Paragraph com marcador (bullet)
EX_HEADING = 89     # Heading 3 "Estatística da Catequese"
EX_ASSINATURA = 95  # centrado


def set_text(el, text):
    """Deixa o parágrafo com um único run (formatação do 1º run) contendo `text`."""
    runs = el.findall(qn("w:r"))
    rpr = None
    if runs:
        first_rpr = runs[0].find(qn("w:rPr"))
        if first_rpr is not None:
            rpr = copy.deepcopy(first_rpr)
    for child in list(el):
        if child.tag != qn("w:pPr"):
            el.remove(child)
    if not text:
        return el
    r = el.makeelement(qn("w:r"), {})
    if rpr is not None:
        r.append(rpr)
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    el.append(r)
    return el


def clone_p(src, text):
    return set_text(copy.deepcopy(src._p if hasattr(src, "_p") else src), text)


def keep_next(el):
    """Mantém o parágrafo na mesma página que o seguinte (ex.: título + tabela)."""
    ppr = el.find(qn("w:pPr"))
    if ppr is None:
        ppr = el.makeelement(qn("w:pPr"), {})
        el.insert(0, ppr)
    if ppr.find(qn("w:keepNext")) is None:
        k = ppr.makeelement(qn("w:keepNext"), {})
        style = ppr.find(qn("w:pStyle"))
        if style is not None:
            style.addnext(k)
        else:
            ppr.insert(0, k)
    return el


def set_cell(tc, text, bold=None):
    ps = tc.findall(qn("w:p"))
    for extra in ps[1:]:
        tc.remove(extra)
    p = ps[0]
    set_text(p, text)
    if bold is not None:
        for r in p.findall(qn("w:r")):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = r.makeelement(qn("w:rPr"), {})
                r.insert(0, rpr)
            for b in rpr.findall(qn("w:b")) + rpr.findall(qn("w:bCs")):
                rpr.remove(b)
            if bold:
                rpr.insert(0, rpr.makeelement(qn("w:b"), {}))


def set_tc_width(tc, w):
    tcw = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
    tcw.set(qn("w:w"), str(w))


def build_table(src_tbl, widths, title, headers, cells, loop):
    """
    Clona a tabela de estatística com `len(widths)` colunas.
      title   — texto da 1ª linha (sobre as colunas 2..n) ou None para omitir
      headers — rótulos da linha de cabeçalho
      cells   — marcadores de cada célula da linha-modelo
      loop    — expressão do {%tr for ... %}
    """
    tbl = copy.deepcopy(src_tbl)
    n = len(widths)
    rows = tbl.findall(qn("w:tr"))
    title_row, head_row, data_row = rows[0], rows[1], rows[2]
    for r in rows[3:]:
        tbl.remove(r)

    # grelha
    grid = tbl.find(qn("w:tblGrid"))
    for g in list(grid):
        grid.remove(g)
    for w in widths:
        g = grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(w)})
        grid.append(g)

    def resize(tr):
        tcs = tr.findall(qn("w:tc"))
        while len(tcs) < n:
            new = copy.deepcopy(tcs[-1])
            tcs[-1].addnext(new)
            tcs.append(new)
        while len(tcs) > n:
            tr.remove(tcs.pop())
        for tc, w in zip(tcs, widths):
            set_tc_width(tc, w)
        return tcs

    # linha de título (ano) — 1 célula vazia + 1 célula com gridSpan n-1
    if title is None:
        tbl.remove(title_row)
    else:
        tcs = title_row.findall(qn("w:tc"))
        set_tc_width(tcs[0], widths[0])
        set_tc_width(tcs[1], sum(widths[1:]))
        span = tcs[1].find(qn("w:tcPr")).find(qn("w:gridSpan"))
        span.set(qn("w:val"), str(n - 1))
        set_cell(tcs[1], title)

    for tc, label in zip(resize(head_row), headers):
        set_cell(tc, label, bold=True)

    for tc, mark in zip(resize(data_row), cells):
        set_cell(tc, mark, bold=False)

    # linhas {%tr for %} / {%tr endfor %}
    open_row = copy.deepcopy(data_row)
    close_row = copy.deepcopy(data_row)
    for i, tc in enumerate(open_row.findall(qn("w:tc"))):
        set_cell(tc, "{%%tr for %s %%}" % loop if i == 0 else "")
    for i, tc in enumerate(close_row.findall(qn("w:tc"))):
        set_cell(tc, "{%tr endfor %}" if i == 0 else "")
    data_row.addprevious(open_row)
    data_row.addnext(close_row)
    return tbl


def main(src_path):
    doc = Document(src_path)
    P = doc.paragraphs
    ex = {
        "titulo": P[EX_TITULO]._p, "lema": P[EX_LEMA]._p, "corpo": P[EX_CORPO]._p,
        "mes": P[EX_MES]._p, "marcador": P[EX_MARCADOR]._p, "heading": P[EX_HEADING]._p,
        "assinatura": P[EX_ASSINATURA]._p,
    }
    ex = {k: copy.deepcopy(v) for k, v in ex.items()}
    src_tbl = copy.deepcopy(doc.tables[0]._tbl)

    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect:
            body.remove(child)

    out = []

    def p(kind, text="", keep=False):
        el = clone_p(ex[kind], text)
        out.append(keep_next(el) if keep else el)

    def tag(text):
        p("corpo", text)

    def bullets(var):
        tag("{%%p for it in %s %%}" % var)
        p("marcador", "{{ it }}")
        tag("{%p endfor %}")

    # ── Cabeçalho ──
    p("titulo", "{{ cab.paroquia }}")
    p("titulo", "{{ cab.comunidades }}")
    p("titulo", "{{ cab.ministerio }}")
    p("corpo", "")
    p("titulo", "Relatório anual das actividades de {{ ano }}")
    tag("{%p if lema %}")
    p("lema", "“{{ lema }}”")
    tag("{%p endif %}")

    # ── Introdução ──
    tag("{%p for par in introducao %}")
    p("corpo", "{{ par }}")
    tag("{%p endfor %}")
    p("corpo", "Destacamos as seguintes actividades:")

    # ── Meses ──
    tag("{%p for m in meses %}")
    p("mes", "{{ m.nome }}")
    tag("{%p for it in m.itens %}")
    p("marcador", "{{ it }}")
    tag("{%p endfor %}")
    tag("{%p endfor %}")

    # ── Ao longo do ano ──
    tag("{%p if ao_longo_do_ano %}")
    p("mes", "Ao longo do Ano {{ ano }}")
    bullets("ao_longo_do_ano")
    tag("{%p endif %}")

    # ── Extraordinárias ──
    tag("{%p if extraordinarias %}")
    p("mes", "Actividades Extraordinárias")
    bullets("extraordinarias")
    tag("{%p endif %}")

    # ── Desafios e Propostas ──
    tag("{%p if desafios or propostas %}")
    p("mes", "Desafios e Propostas:")
    tag("{%p endif %}")
    tag("{%p if desafios %}")
    p("mes", "Desafios")
    bullets("desafios")
    tag("{%p endif %}")
    tag("{%p if propostas %}")
    p("mes", "Propostas")
    bullets("propostas")
    tag("{%p endif %}")

    # ── Estatística ──
    p("heading", "Estatística da Catequese")
    tag("{%p for par in estatistica_paragrafos %}")
    p("corpo", "{{ par }}")
    tag("{%p endfor %}")
    out.append(build_table(
        src_tbl,
        widths=[2900, 1600, 1600, 1500, 1753],
        title="{{ ano }}",
        headers=["Comunidades", "Santa Ana", "Assunção", "Total", "Assunção {{ ano_anterior }}"],
        cells=["{{ r.indicador }}", "{{ r.santa_ana }}", "{{ r.sede }}", "{{ r.total }}", "{{ r.ano_anterior }}"],
        loop="r in estatistica",
    ))
    tag("{%p for n in notas_estatistica %}")
    p("corpo", "{{ n }}")
    tag("{%p endfor %}")

    # ── Catequizandos por fase ──
    tag("{%p if fases %}")
    p("corpo", "")
    p("corpo", "Distribuição dos catequizandos da comunidade sede por fase:", keep=True)
    out.append(build_table(
        src_tbl,
        widths=[3953, 1600, 2000, 1800],
        title=None,
        headers=["Fase", "Turmas", "Catequizandos", "Desistências"],
        cells=["{{ f.fase }}", "{{ f.turmas }}", "{{ f.catequizandos }}", "{{ f.desistencias }}"],
        loop="f in fases",
    ))
    tag("{%p endif %}")

    # ── Finanças (opcional) ──
    tag("{%p if financas %}")
    p("corpo", "")
    p("heading", "Resumo Financeiro", keep=True)
    p("corpo", "{{ financas.texto }}", keep=True)
    out.append(build_table(
        src_tbl,
        widths=[4553, 2400, 2400],
        title=None,
        headers=["Rubrica", "Receitas (MT)", "Despesas (MT)"],
        cells=["{{ x.rubrica }}", "{{ x.receita }}", "{{ x.despesa }}"],
        loop="x in financas.linhas",
    ))
    tag("{%p endif %}")

    # ── Assinatura ──
    p("corpo", "")
    p("corpo", "")
    p("assinatura", "A Coordenação da Catequese e Formação Permanente")
    p("corpo", "")
    p("assinatura", "{{ local_data }}")

    for el in out:
        sect.addprevious(el)

    doc.core_properties.title = "Relatório Anual — Ministério da Catequese"
    doc.save(OUT)
    print("Modelo gravado em", OUT)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "docs/Relatório 2025.docx")
